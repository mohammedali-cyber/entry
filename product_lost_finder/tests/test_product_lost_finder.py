import io

from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests import TransactionCase, new_test_user, tagged
from odoo.tools import mute_logger

try:
    import openpyxl
except ImportError:
    openpyxl = None


@tagged('post_install', '-at_install')
class TestProductLostFinder(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.wh_main = cls.env['stock.warehouse'].search([('company_id', '=', cls.company.id)], limit=1)
        cls.wh_branch = cls.env['stock.warehouse'].create({
            'name': 'PLF Branch Warehouse',
            'code': 'PLFB',
            'company_id': cls.company.id,
        })

        size = cls.env['product.attribute'].create({'name': 'PLF Size'})
        v38, v39 = cls.env['product.attribute.value'].create([
            {'name': '38', 'attribute_id': size.id},
            {'name': '39', 'attribute_id': size.id},
        ])
        cls.shoe = cls.env['product.template'].create({
            'name': 'PLF Shoe',
            'is_storable': True,
            'list_price': 29.0,
            'attribute_line_ids': [(0, 0, {'attribute_id': size.id, 'value_ids': [(6, 0, [v38.id, v39.id])]})],
        })
        cls.shoe_38 = cls.shoe.product_variant_ids.filtered(
            lambda p: p.product_template_attribute_value_ids.product_attribute_value_id == v38)
        cls.shoe_39 = cls.shoe.product_variant_ids - cls.shoe_38
        # Like real data: every variant has its own reference built on the model number.
        cls.shoe_38.write({'default_code': 'PLF100-38', 'barcode': '9990038'})
        cls.shoe_39.default_code = 'PLF100-39'
        cls.pencil = cls.env['product.product'].create({
            'name': 'PLF Pencil', 'is_storable': True, 'default_code': 'PLF-001', 'barcode': '9990001'})
        cls.service = cls.env['product.product'].create({'name': 'PLF Delivery Service', 'type': 'service'})

        Quant = cls.env['stock.quant']
        # Shoe 38: none in main branch, 4 in the other branch. Shoe 39: 3 in main branch.
        Quant._update_available_quantity(cls.shoe_38, cls.wh_branch.lot_stock_id, 4)
        Quant._update_available_quantity(cls.shoe_39, cls.wh_main.lot_stock_id, 3)
        Quant._update_available_quantity(cls.pencil, cls.wh_main.lot_stock_id, -2)

        cls.user_main = new_test_user(cls.env, login='plf_user_main', groups='stock.group_stock_user')
        cls.user_main.lost_finder_warehouse_id = cls.wh_main
        cls.user_branch = new_test_user(cls.env, login='plf_user_branch', groups='stock.group_stock_user')
        cls.user_branch.lost_finder_warehouse_id = cls.wh_branch
        cls.user_no_branch = new_test_user(cls.env, login='plf_user_nobranch', groups='stock.group_stock_user')
        cls.manager = new_test_user(cls.env, login='plf_manager',
                                    groups='product_lost_finder.group_lost_finder_manager')
        cls.basic_user = new_test_user(cls.env, login='plf_basic_user', groups='base.group_user')

    def _finder(self, user):
        return self.env['product.lost.finder'].with_user(user)

    def _rows(self, user, template, warehouse):
        result = self._finder(user).get_product_tables([template.id], warehouse.id)
        return result, {row['id']: row for row in result['tables'][0]['variants']}

    # Search ------------------------------------------------------------

    def test_search_templates(self):
        finder = self._finder(self.user_main)
        for term in ('PLF Shoe', 'plf100', '9990038', '99900'):
            ids = [t['id'] for t in finder.search_templates(term)['templates']]
            self.assertIn(self.shoe.id, ids, term)
        self.assertEqual([t['id'] for t in finder.search_templates('9990001')['templates']],
                         [self.pencil.product_tmpl_id.id])
        ids = [t['id'] for t in finder.search_templates('PLF')['templates']]
        self.assertNotIn(self.service.product_tmpl_id.id, ids, "Services are excluded")
        self.assertEqual(finder.search_templates('PLF-not-existing'), {'templates': [], 'total': 0})
        self.assertEqual(finder.search_templates('  '), {'templates': [], 'total': 0})

    # Product table -----------------------------------------------------

    def test_table_quantities_per_branch(self):
        result, rows = self._rows(self.user_main, self.shoe, self.wh_main)
        self.assertEqual(result['warehouses'][0]['id'], self.wh_main.id, "User's branch is the first column")
        main, branch = str(self.wh_main.id), str(self.wh_branch.id)
        self.assertEqual(rows[self.shoe_38.id]['qty_by_warehouse'][main], 0)
        self.assertEqual(rows[self.shoe_38.id]['qty_by_warehouse'][branch], 4)
        self.assertEqual(rows[self.shoe_38.id]['total'], 4)
        self.assertTrue(rows[self.shoe_38.id]['is_lost'])
        self.assertFalse(rows[self.shoe_39.id]['is_lost'])
        self.assertEqual(rows[self.shoe_38.id]['name'], '38')
        variants = result['tables'][0]['variants']
        self.assertEqual(variants[0]['id'], self.shoe_39.id, "Available variants come first")

    def test_negative_quantity_is_lost(self):
        _result, rows = self._rows(self.user_main, self.pencil.product_tmpl_id, self.wh_main)
        self.assertTrue(rows[self.pencil.id]['is_lost'])

    def test_employee_is_bound_to_own_branch(self):
        # A regular employee asking for another branch still gets their own branch.
        result, rows = self._rows(self.user_branch, self.shoe, self.wh_main)
        self.assertEqual(result['branch_id'], self.wh_branch.id)
        self.assertFalse(rows[self.shoe_38.id]['is_lost'])
        with self.assertRaises(UserError):
            self._finder(self.user_no_branch).get_product_tables([self.shoe.id], self.wh_main.id)
        self.assertFalse(self._finder(self.user_no_branch).get_context()['branch_id'])

    def test_manager_chooses_branch(self):
        ctx = self._finder(self.manager).get_context()
        self.assertTrue(ctx['is_manager'])
        result, rows = self._rows(self.manager, self.shoe, self.wh_branch)
        self.assertEqual(result['branch_id'], self.wh_branch.id)
        self.assertFalse(rows[self.shoe_38.id]['is_lost'])

    def test_access_rights(self):
        finder = self._finder(self.basic_user)
        with self.assertRaises(AccessError):
            finder.get_context()
        with self.assertRaises(AccessError):
            finder.search_templates('PLF')
        with self.assertRaises(AccessError):
            finder.get_product_tables([self.shoe.id], self.wh_main.id)
        with self.assertRaises(AccessError):
            finder.export_xlsx([self.shoe_38.id], self.wh_main.id)

    # Lost sales --------------------------------------------------------

    def test_record_lost_sale(self):
        LostSale = self.env['product.lost.sale'].with_user(self.user_main)
        sale = LostSale.browse(LostSale.record_lost_sale(self.shoe_38.id, self.wh_main.id, 2, 'Size 38'))
        self.assertEqual(sale.warehouse_id, self.wh_main)
        self.assertEqual(sale.user_id, self.user_main)
        self.assertEqual(sale.quantity, 2)
        self.assertEqual(sale.price_unit, self.shoe_38.lst_price)
        self.assertEqual(sale.amount, 2 * self.shoe_38.lst_price)
        self.assertEqual(sale.note, 'Size 38')

    def test_record_lost_sale_refused_when_available(self):
        LostSale = self.env['product.lost.sale'].with_user(self.user_main)
        with self.assertRaises(UserError):
            LostSale.record_lost_sale(self.shoe_39.id, self.wh_main.id, 1)
        with self.assertRaises(ValidationError), mute_logger('odoo.sql_db'):
            LostSale.record_lost_sale(self.shoe_38.id, self.wh_main.id, 0)
        with self.assertRaises(AccessError):
            self.env['product.lost.sale'].with_user(self.basic_user).record_lost_sale(
                self.shoe_38.id, self.wh_main.id, 1)

    def test_lost_sales_visible_per_branch(self):
        self.env['product.lost.sale'].with_user(self.user_main).record_lost_sale(self.shoe_38.id, self.wh_main.id, 1)
        self.env['product.lost.sale'].with_user(self.user_branch).record_lost_sale(self.shoe_39.id, self.wh_branch.id, 1)
        domain = [('product_tmpl_id', '=', self.shoe.id)]
        mine = self.env['product.lost.sale'].with_user(self.user_main).search(domain)
        self.assertEqual(mine.warehouse_id, self.wh_main, "Employees only see their own branch")
        self.assertFalse(self.env['product.lost.sale'].with_user(self.user_no_branch).search(domain))
        everything = self.env['product.lost.sale'].with_user(self.manager).search(domain)
        self.assertEqual(everything.warehouse_id, self.wh_main | self.wh_branch)
        with self.assertRaises(AccessError):
            mine.with_user(self.user_main).unlink()

    # Excel -------------------------------------------------------------

    def test_export_xlsx(self):
        finder = self._finder(self.user_main).with_context(lang='en_US')
        filename, content = finder.export_xlsx([self.shoe_38.id, self.shoe_39.id], self.wh_main.id)
        self.assertTrue(filename.endswith('.xlsx'))
        self.assertEqual(content[:2], b'PK')
        if not openpyxl:
            return
        sheet = openpyxl.load_workbook(io.BytesIO(content)).active
        rows = list(sheet.iter_rows(min_row=5, values_only=True))
        self.assertEqual(len(rows), 2)
        code, name, variant, barcode, price, qty_main, qty_branch, total, status = rows[0]
        self.assertEqual((name, variant, barcode), ('PLF Shoe', '38', '9990038'))
        self.assertEqual(price, self.shoe_38.lst_price)
        self.assertEqual((qty_main, qty_branch, total), (0, 4, 4))
        self.assertEqual(status, 'Lost')
        self.assertEqual(rows[1][-1], 'Available')
        with self.assertRaises(UserError):
            finder.export_xlsx([], self.wh_main.id)


@tagged('post_install', '-at_install')
class TestLostSaleDetailReport(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.wh_main = cls.env['stock.warehouse'].search([('company_id', '=', cls.env.company.id)], limit=1)
        cls.wh_branch = cls.env['stock.warehouse'].create({'name': 'PLF Detail Branch', 'code': 'PLFD'})
        Attribute, Value = cls.env['product.attribute'], cls.env['product.attribute.value']
        cls.color = Attribute.create({'name': 'PLF color', 'sequence': 1})
        cls.size = Attribute.create({'name': 'PLF size', 'sequence': 2})
        cls.white, cls.silver = Value.create([{'name': 'WHITE', 'attribute_id': cls.color.id, 'sequence': 1},
                                              {'name': 'SILVER', 'attribute_id': cls.color.id, 'sequence': 2}])
        cls.s38, cls.s39 = Value.create([{'name': '38', 'attribute_id': cls.size.id, 'sequence': 1},
                                         {'name': '39', 'attribute_id': cls.size.id, 'sequence': 2}])

        def template(name, colors, sizes):
            return cls.env['product.template'].create({
                'name': name, 'is_storable': True, 'list_price': 10.0,
                'attribute_line_ids': [
                    (0, 0, {'attribute_id': cls.color.id, 'value_ids': [(6, 0, colors.ids)]}),
                    (0, 0, {'attribute_id': cls.size.id, 'value_ids': [(6, 0, sizes.ids)]}),
                ],
            })
        cls.shoe = template('PLF Detail Shoe', cls.white | cls.silver, cls.s38 | cls.s39)
        cls.sandal = template('PLF Detail Sandal', cls.white, cls.s38)

        def variant(tmpl, *values):
            return tmpl.product_variant_ids.filtered(
                lambda p: p.product_template_attribute_value_ids.product_attribute_value_id == cls.env[
                    'product.attribute.value'].concat(*values))

        cls.shoe_w38 = variant(cls.shoe, cls.white, cls.s38)
        cls.shoe_s39 = variant(cls.shoe, cls.silver, cls.s39)
        cls.sandal_w38 = variant(cls.sandal, cls.white, cls.s38)

        cls.manager = new_test_user(cls.env, login='plf_detail_manager',
                                    groups='product_lost_finder.group_lost_finder_manager', tz='UTC')
        cls.employee = new_test_user(cls.env, login='plf_detail_employee', groups='stock.group_stock_user', tz='UTC')
        cls.employee.lost_finder_warehouse_id = cls.wh_main

        LostSale = cls.env['product.lost.sale']
        LostSale.create([
            {'product_id': cls.shoe_w38.id, 'warehouse_id': cls.wh_main.id, 'quantity': 5, 'price_unit': 10,
             'date': '2026-10-05 10:00:00'},
            {'product_id': cls.shoe_s39.id, 'warehouse_id': cls.wh_main.id, 'quantity': 1, 'price_unit': 10,
             'date': '2026-10-06 10:00:00'},
            {'product_id': cls.sandal_w38.id, 'warehouse_id': cls.wh_branch.id, 'quantity': 2, 'price_unit': 10,
             'date': '2026-09-15 10:00:00'},
        ])

    def _report(self, user, **filters):
        filters.setdefault('column_attribute_id', self.color.id)
        return self.env['product.lost.sale'].with_user(user).get_detail_report(filters)

    def _matrix(self, report):
        names = {c['key']: c['name'] for c in report['columns']}
        return {row['label']: {names[k]: q for k, q in row['cells'].items()} for row in report['rows']}

    def test_options(self):
        options = self.env['product.lost.sale'].with_user(self.manager).get_detail_options()
        attributes = {a['name']: [v['name'] for v in a['values']] for a in options['attributes']}
        self.assertEqual(attributes['PLF color'], ['WHITE', 'SILVER'])
        self.assertEqual(attributes['PLF size'], ['38', '39'])
        self.assertEqual(options['column_attribute_id'], self.color.id, "Color attribute is the default column")
        names = [t['name'] for t in options['templates']]
        self.assertIn('PLF Detail Shoe', names)
        self.assertIn('PLF Detail Sandal', names)

    def test_product_filter_shows_colors_without_sizes(self):
        report = self._report(self.manager, template_id=self.shoe.id)
        self.assertEqual(self._matrix(report), {'PLF Detail Shoe': {'WHITE': 5, 'SILVER': 1}})
        self.assertEqual(report['total'], 6)
        self.assertEqual(report['amount'], 60)
        self.assertEqual([c['name'] for c in report['columns']], ['WHITE', 'SILVER'])

    def test_size_only_narrows_when_chosen(self):
        report = self._report(self.manager, template_id=self.shoe.id, value_ids=[self.s39.id])
        self.assertEqual(self._matrix(report), {'PLF Detail Shoe': {'SILVER': 1}})

    def test_color_filter_shows_all_products_with_that_color(self):
        report = self._report(self.manager, value_ids=[self.white.id])
        self.assertEqual(self._matrix(report), {'PLF Detail Shoe': {'WHITE': 5}, 'PLF Detail Sandal': {'WHITE': 2}})
        report = self._report(self.manager, value_ids=[self.white.id, self.s39.id])
        self.assertFalse(report['rows'])

    def test_branch_and_period_filters(self):
        report = self._report(self.manager, warehouse_id=self.wh_branch.id)
        self.assertEqual(self._matrix(report), {'PLF Detail Sandal': {'WHITE': 2}})
        report = self._report(self.manager, date_from='2026-10-01', date_to='2026-10-05')
        self.assertEqual(self._matrix(report), {'PLF Detail Shoe': {'WHITE': 5}})

    def test_employee_sees_own_branch_only(self):
        report = self._report(self.employee, warehouse_id=self.wh_branch.id)
        self.assertEqual(self._matrix(report), {'PLF Detail Shoe': {'WHITE': 5, 'SILVER': 1}})
        options = self.env['product.lost.sale'].with_user(self.employee).get_detail_options()
        self.assertEqual([w['id'] for w in options['warehouses']], [self.wh_main.id])
        self.assertNotIn('PLF Detail Sandal', [t['name'] for t in options['templates']])

    def test_export_detail_xlsx(self):
        LostSale = self.env['product.lost.sale'].with_user(self.manager).with_context(lang='en_US')
        filename, content = LostSale.export_detail_xlsx({'template_id': self.shoe.id,
                                                         'column_attribute_id': self.color.id})
        self.assertTrue(filename.endswith('.xlsx'))
        if not openpyxl:
            return
        sheet = openpyxl.load_workbook(io.BytesIO(content)).active
        rows = list(sheet.iter_rows(min_row=4, values_only=True))
        self.assertEqual(rows[0], ('Product', 'WHITE', 'SILVER', 'Total', 'Lost Amount'))
        self.assertEqual(rows[1], ('PLF Detail Shoe', 5, 1, 6, 60))
        self.assertEqual(rows[-1], ('Total', 5, 1, 6, 60))
